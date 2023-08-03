from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import (
    coordinate_to_tuple,
    column_index_from_string,
    cols_from_range,
    get_column_letter,
)
from openpyxl.cell import MergedCell
import yaml

config_filepath = "domeo210_modbus.yml"
data = yaml.load(open(config_filepath, "r"), yaml.SafeLoader)
wb = load_workbook("Downloads/DOMEO-PARAMETRES-MODBUS.xlsx")

ws = wb.active

headers = dict(
    register_number='A',
    description='B',
    data_type='C',
    data_type_value='D',
    mode='E',
    default='F',
    comments='G'
)
reverse_headers = {v: k for k, v in headers.items()}
tables = dict(discrete_inputs=ws["A6":"G25"],
              coils=ws["A28":"G45"],
              input_registers=ws["A48":"G97"],
              holding_registers=ws["A100":"G121"])

data = dict(discrete_inputs={}, coils={}, input_registers={}, holding_registers={})

for name, table in tables.items():
    data[name] = []
    for line_index, lines in enumerate(table):
        values = dict(register_number='',
                      description='',
                      data=[],
                      mode='',
                      default='',
                      comments='')
        cache = False
        register_number = {'count': 0, 'value': ''}
        for index, cell in enumerate(lines):
            if isinstance(cell, MergedCell):
                continue
            value = ws[f"{cell.column_letter}{cell.row}"].value
            if reverse_headers[cell.column_letter] == 'data_type':
                register_number['count'] = value
            elif reverse_headers[cell.column_letter] == 'data_type_value':
                register_number['value'] = value
                values['data'].append(register_number)
            else:
                if values[reverse_headers[cell.column_letter]] == '':
                    values[reverse_headers[cell.column_letter]] = value
        if values['register_number'] in ('', None):
            data[name][-1]['data'].append(register_number)
        else:
            data[name].append(values)

with open(config_filepath, "w") as f:
    f.write(yaml.dump(data))
